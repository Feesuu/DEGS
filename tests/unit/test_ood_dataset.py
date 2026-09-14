from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

import degs.ood_dataset as ood_dataset


def _input_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Table"
    sheet.append(["City", "Value"])
    sheet.append(["A", 7])
    ood_dataset._style_workbook(workbook)
    return ood_dataset._workbook_bytes(workbook)


def test_input_population_exposes_question_but_never_gold(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setitem(
        ood_dataset.SOURCE_SPECS,
        "wikitq",
        {
            **ood_dataset.SOURCE_SPECS["wikitq"],
            "count": 1,
        },
    )
    manifest = ood_dataset._publish_population(
        dataset="wikitq",
        rows=[("dev-1", "wikitq_dev-1", "What is the value for A?", _input_workbook())],
        output=tmp_path,
    )
    verified = ood_dataset.verify_population(tmp_path)
    assert verified == manifest
    assert verified["gold_available"] is False
    task = verified["tasks"][0]
    assert task["semantic_query"] == "What is the value for A?"
    dataset_bytes = (tmp_path / "dataset.json").read_bytes()
    assert b"gold" not in dataset_bytes.lower()
    assert b"targetvalue" not in dataset_bytes.lower()
    rows = json.loads(dataset_bytes)
    assert rows[0]["question"] == task["semantic_query"]
    workbook = load_workbook(tmp_path / task["spreadsheet_path"] / task["input_file"])
    try:
        assert workbook["Answer"]["B1"].value is None
    finally:
        workbook.close()


def test_safe_relative_rejects_parent_escape() -> None:
    try:
        ood_dataset._safe_relative("../gold.xlsx", label="path")
    except ValueError as exc:
        assert "safe relative path" in str(exc)
    else:
        raise AssertionError("parent traversal was accepted")


def test_population_rejects_dataset_routing_away_from_manifest(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setitem(
        ood_dataset.SOURCE_SPECS,
        "wikitq",
        {**ood_dataset.SOURCE_SPECS["wikitq"], "count": 1},
    )
    ood_dataset._publish_population(
        dataset="wikitq",
        rows=[("dev-1", "wikitq_dev-1", "Question?", _input_workbook())],
        output=tmp_path,
    )
    rows = json.loads((tmp_path / "dataset.json").read_text(encoding="utf-8"))
    rows[0]["spreadsheet_path"] = "spreadsheet/other"
    dataset_bytes = json.dumps(rows, ensure_ascii=False, indent=2).encode("utf-8") + b"\n"
    (tmp_path / "dataset.json").write_bytes(dataset_bytes)
    manifest_path = tmp_path / "retrieval_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["dataset_json_sha256"] = ood_dataset._sha(dataset_bytes)
    unsigned = {key: value for key, value in manifest.items() if key != "self_sha256"}
    manifest["self_sha256"] = ood_dataset._sha(
        ood_dataset.canonical_json_bytes(unsigned)
    )
    manifest_path.write_bytes(ood_dataset.canonical_json_bytes(manifest))

    with pytest.raises(ValueError, match="prepared task"):
        ood_dataset.verify_population(tmp_path)
