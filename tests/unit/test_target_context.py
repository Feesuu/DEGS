from __future__ import annotations

import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
import pytest

from degs.experience_simgrag import need_retrieval_document
from degs.section_graph import IOContract
from degs.target_context import (
    TARGET_EVIDENCE_POLICY,
    TargetContextUnavailable,
    TargetEvidenceCard,
    build_target_evidence_card,
)
from degs.retrieval_clarification import (
    clarification_response_schema,
    parse_retrieval_clarification,
)
from degs.workflow_retrieval import (
    NeedGraph,
    NeedNode,
)


def _dataset(tmp_path: Path) -> tuple[Path, Path]:
    dataset = tmp_path / "dataset.json"
    dataset.write_text("[]\n", encoding="utf-8")
    task_dir = tmp_path / "spreadsheet" / "task"
    task_dir.mkdir(parents=True)
    return dataset, task_dir


def _workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    target = workbook.active
    target.title = "Report"
    target["A1"] = "Day"
    target["B1"] = "Month"
    target["C1"] = "Year"
    target["J1"] = "Date"
    target["A2"] = 1
    target["B2"] = 2
    target["C2"] = 2024
    target["J2"] = "=DATE(C2,B2,A2)"
    target["J2"].number_format = "yyyy-mm-dd"
    source = workbook.create_sheet("Data")
    source["A1"] = "Account"
    source["B1"] = "Amount"
    source["A2"] = "A"
    source["B2"] = 10
    workbook.save(path)


def _need_graph() -> NeedGraph:
    contract = (IOContract("value", "a spreadsheet value"),)
    output = (IOContract("value", "the target value representation"),)
    return NeedGraph(
        (
            NeedNode(
                "Construct the target value",
                ("The target representation must match the workbook.",),
                contract,
                output,
            ),
        ),
        (),
    )


def test_target_evidence_card_observes_relevant_sheets_and_target_landmarks(
    tmp_path: Path,
) -> None:
    dataset, task_dir = _dataset(tmp_path)
    _workbook(task_dir / "1_task_init.xlsx")
    card = build_target_evidence_card(
        dataset_path=dataset,
        spreadsheet_path="spreadsheet/task",
        answer_position="'Report'!J2:J8",
        instruction="Use Report values and look up supporting data from the Data sheet.",
    )
    assert card.status == "OK"
    assert card.audit_payload()["policy"] == TARGET_EVIDENCE_POLICY
    assert all(
        row["kind"] != "declared_answer_position"
        for row in card.llm_payload()["input_observations"]
    )
    observations = {
        row["kind"]: [] for row in card.llm_payload()["input_observations"]
    }
    for row in card.llm_payload()["input_observations"]:
        observations[row["kind"]].append(row["value"])
    assert observations["workbook_sheet_inventory"] == [
        {"active_sheet": "Report", "sheet_names": ["Report", "Data"]}
    ]
    structures = {
        row["sheet"]: row["rows"]
        for row in observations["relevant_sheet_structure"]
    }
    assert set(structures) == {"Report", "Data"}
    assert "range" not in observations["target_range_landmarks"][0]
    landmarks = observations["target_range_landmarks"][0]["cells"]
    j2 = next(row for row in landmarks if row["coordinate"] == "J2")
    assert j2["value"] == "<formula>"
    assert j2["data_type"] == "f"
    assert j2["number_format"] == "yyyy-mm-dd"


def test_formula_evidence_is_content_free_and_deterministic(tmp_path: Path) -> None:
    dataset, task_dir = _dataset(tmp_path)
    _workbook(task_dir / "1_task_init.xlsx")
    kwargs = {
        "dataset_path": dataset,
        "spreadsheet_path": "spreadsheet/task",
        "answer_position": "'Report'!J2:J8",
        "instruction": "Fill Report column J.",
    }
    first = build_target_evidence_card(**kwargs)
    second = build_target_evidence_card(**kwargs)
    assert first.llm_payload() == second.llm_payload()
    assert first.audit_payload() == second.audit_payload()
    payload = json.dumps(first.llm_payload(), ensure_ascii=False)
    assert "=DATE(" not in payload
    assert "0x" not in payload


def test_workflow_context_keeps_roles_but_masks_non_header_values(
    tmp_path: Path,
) -> None:
    dataset, task_dir = _dataset(tmp_path)
    _workbook(task_dir / "1_task_init.xlsx")
    card = build_target_evidence_card(
        dataset_path=dataset,
        spreadsheet_path="spreadsheet/task",
        answer_position="'Report'!J2:J8",
        instruction="Fill the Date result in Report column J.",
    )
    evidence_ids = tuple(
        row["evidence_id"]
        for row in card.observations
        if row["kind"] in {"relevant_sheet_structure", "target_range_landmarks"}
    )
    document = card.workflow_context_document(
        "Fill the Date result in Report column J.",
        evidence_ids=evidence_ids,
    )
    assert document.startswith("Task instruction:")
    assert 'J1="Date"' in document
    assert "J2=<formula>" in document
    assert "=DATE(" not in document
    assert "2024" not in document
    assert "J2:J8" not in document
    with pytest.raises(ValueError, match="differs"):
        card.workflow_context_document(
            "Fill the Date result in Report column J.",
            evidence_ids=("O999",),
        )


def test_unavailable_workflow_context_falls_back_to_query_only() -> None:
    card = TargetEvidenceCard(
        status="UNAVAILABLE",
        observations=(),
        input_sha256=None,
        input_relative_path=None,
        failure="input unavailable",
    )
    assert card.workflow_context_document("Inspect the workbook.") == (
        "Inspect the workbook."
    )


def test_workflow_context_does_not_guess_that_first_nonempty_row_is_header(
    tmp_path: Path,
) -> None:
    dataset, task_dir = _dataset(tmp_path)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A2"] = "Private Customer Name"
    sheet["A3"] = "Another Customer"
    workbook.save(task_dir / "1_task_init.xlsx")
    card = build_target_evidence_card(
        dataset_path=dataset,
        spreadsheet_path="spreadsheet/task",
        answer_position="A4",
        instruction="Fill the result in column A.",
    )
    structure_ids = tuple(
        row["evidence_id"]
        for row in card.observations
        if row["kind"] == "relevant_sheet_structure"
    )
    document = card.workflow_context_document(
        "Fill the result in column A.", evidence_ids=structure_ids
    )
    assert "Private Customer Name" not in document
    assert "Another Customer" not in document
    assert "A2=<text>" in document


def test_workflow_context_run_length_encodes_wide_landmarks_without_truncation() -> None:
    cells = [
        {
            "coordinate": f"{openpyxl.utils.get_column_letter(column)}1",
            "value": column,
            "data_type": "n",
            "number_format": "0",
        }
        for column in range(1, 1002)
    ]
    card = TargetEvidenceCard(
        status="OK",
        observations=(
            {
                "evidence_id": "O0",
                "kind": "target_range_landmarks",
                "content": json.dumps(
                    {"sheet": "Sheet1", "cells": cells, "merged_ranges": []}
                ),
            },
        ),
        input_sha256="a" * 64,
        input_relative_path="spreadsheet/task/input.xlsx",
    )

    document = card.workflow_context_document("Transform the wide table")

    assert "A1:ALM1=<numeric_or_blank>[numeric_or_other]" in document
    assert len(document) < 300


def test_wide_sheet_structure_is_a_semantic_projection_not_a_cell_dump(
    tmp_path: Path,
) -> None:
    dataset, task_dir = _dataset(tmp_path)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Wide"
    for column in range(1, 401):
        sheet.cell(row=1, column=column, value=f"Header {column}")
        sheet.cell(row=2, column=column, value=column)
    workbook.save(task_dir / "1_task_init.xlsx")

    card = build_target_evidence_card(
        dataset_path=dataset,
        spreadsheet_path="spreadsheet/task",
        answer_position="GK3",
        instruction="Use the value in column GK and place the result in GK3.",
    )
    structure = next(
        row["value"]
        for row in card.llm_payload()["input_observations"]
        if row["kind"] == "relevant_sheet_structure"
    )
    assert structure["rows"][0]["nonempty_column_spans"] == ["A:OJ"]
    assert {
        cell["coordinate"] for cell in structure["rows"][0]["selected_cells"]
    } == {"GJ1", "GK1", "GL1"}
    assert len(json.dumps(card.llm_payload(), ensure_ascii=False)) < 5_000


def test_target_evidence_card_rejects_paths_outside_dataset(tmp_path: Path) -> None:
    dataset, task_dir = _dataset(tmp_path)
    _workbook(task_dir / "1_task_init.xlsx")
    with pytest.raises(ValueError, match="beneath"):
        build_target_evidence_card(
            dataset_path=dataset,
            spreadsheet_path="../outside",
            answer_position="A1",
            instruction="Inspect the workbook.",
        )
    (task_dir / "1_task_golden.xlsx").write_bytes(
        (task_dir / "1_task_init.xlsx").read_bytes()
    )
    card = build_target_evidence_card(
        dataset_path=dataset,
        spreadsheet_path="spreadsheet/task",
        answer_position="A1",
        instruction="Inspect the workbook.",
    )
    assert card.input_relative_path == "spreadsheet/task/1_task_init.xlsx"
    assert all(
        set(row) == {"evidence_id", "kind", "content_sha256"}
        for row in card.audit_payload()["observations"]
    )


def test_target_evidence_ignores_golden_named_input(
    tmp_path: Path,
) -> None:
    dataset, task_dir = _dataset(tmp_path)
    _workbook(task_dir / "answer_golden_init.xlsx")
    with pytest.raises(TargetContextUnavailable, match="unavailable"):
        build_target_evidence_card(
            dataset_path=dataset,
            spreadsheet_path="spreadsheet/task",
            answer_position="A1",
            instruction="Inspect the workbook.",
        )


def test_corrupt_worksheet_xml_is_item_local_unavailable(tmp_path: Path) -> None:
    dataset, task_dir = _dataset(tmp_path)
    good = task_dir / "good.xlsx"
    broken = task_dir / "1_task_init.xlsx"
    _workbook(good)
    with ZipFile(good) as source, ZipFile(broken, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                content = b"<worksheet><broken>"
            target.writestr(info, content)
    good.unlink()
    with pytest.raises(TargetContextUnavailable, match="could not be parsed"):
        build_target_evidence_card(
            dataset_path=dataset,
            spreadsheet_path="spreadsheet/task",
            answer_position="A1",
            instruction="Inspect the workbook.",
        )


def test_retrieval_clarification_is_strict_and_preserves_unmodified_document() -> None:
    graph = _need_graph()
    evidence_ids = ("Q", "O0", "O1")
    schema = clarification_response_schema(graph, evidence_ids=evidence_ids)
    assert schema["properties"]["clarifications"]["items"]["properties"][
        "need_node"
    ]["enum"] == [0]
    keep = parse_retrieval_clarification(
        {"clarifications": []},
        need_graph=graph,
        evidence_ids=evidence_ids,
    )
    assert keep.by_need_node == {}
    assert need_retrieval_document(graph.nodes[0]) == graph.nodes[0].description

    clarified = parse_retrieval_clarification(
        {
            "clarifications": [
                {
                    "need_node": 0,
                    "evidence_ids": ["O1"],
                }
            ],
        },
        need_graph=graph,
        evidence_ids=evidence_ids,
    )
    card = TargetEvidenceCard(
        status="OK",
        observations=(
            {
                "evidence_id": "O1",
                "kind": "target_range_landmarks",
                "content": json.dumps(
                    {
                        "cells": [
                            {
                                "value": "2024-02-01",
                                "data_type": "d",
                                "number_format": "yyyy-mm-dd",
                            }
                        ]
                    }
                ),
            },
        ),
        input_sha256="a" * 64,
        input_relative_path="spreadsheet/task/1_task_init.xlsx",
    )
    document = need_retrieval_document(
        graph.nodes[0],
        card.retrieval_clarification(clarified.by_need_node[0]),
    )
    assert document.startswith(graph.nodes[0].description)
    assert "date_or_time" in document


@pytest.mark.parametrize(
    "value",
    [
        {"clarifications": [{"need_node": 0}]},
        {
            "clarifications": [
                {
                    "need_node": 0,
                    "retrieval_clarification": "A distinction",
                    "evidence_ids": ["Q", "O9"],
                }
            ],
        },
        {
            "clarifications": [
                {
                    "need_node": 0,
                    "retrieval_clarification": "A distinction",
                    "evidence_ids": [["O0"]],
                }
            ],
        },
        {
            "clarifications": [
                {
                    "need_node": 0,
                    "retrieval_clarification": "=SUM(A1:A2); write it into B7",
                    "evidence_ids": ["O0"],
                }
            ],
        },
        {
            "clarifications": [
                {
                    "need_node": 0,
                    "retrieval_clarification": "The exact answer is 42",
                    "evidence_ids": ["O0"],
                }
            ],
        },
        {
            "clarifications": [
                {
                    "need_node": 0,
                    "retrieval_clarification": "A distinction",
                    "evidence_ids": ["Q"],
                }
            ],
        },
    ],
)
def test_retrieval_clarification_rejects_inconsistent_or_uncited_output(
    value: object,
) -> None:
    with pytest.raises(ValueError):
        parse_retrieval_clarification(
            value,
            need_graph=_need_graph(),
            evidence_ids=("Q", "O0", "O1"),
        )
