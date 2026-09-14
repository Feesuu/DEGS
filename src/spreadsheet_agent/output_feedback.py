from __future__ import annotations

import hashlib
import tempfile
from pathlib import Path
from typing import Any, Callable

import openpyxl
from openpyxl.utils.cell import get_column_letter, range_boundaries

from react_agent import Tool


OUTPUT_FEEDBACK_POLICY = "FIRST_REAL_MUTATION_LIBREOFFICE_TARGET_SNAPSHOT_V1"
_MAX_INSPECTED_CELLS = 24
_FORMULA_ERRORS = {
    "#DIV/0!",
    "#N/A",
    "#NAME?",
    "#NULL!",
    "#NUM!",
    "#REF!",
    "#VALUE!",
}


def _sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _display(value: Any) -> str:
    text = repr(value)
    return text if len(text) <= 120 else text[:117] + "..."


def _sample_indices(total: int, maximum: int = _MAX_INSPECTED_CELLS) -> list[int]:
    if total <= maximum:
        return list(range(total))
    return sorted({round(index * (total - 1) / (maximum - 1)) for index in range(maximum)})


def _target_ranges(workbook, answer_position: str):
    for raw in answer_position.split(","):
        item = raw.strip()
        if not item:
            continue
        if "!" in item:
            sheet_name, cell_range = item.rsplit("!", 1)
            sheet_name = sheet_name.strip().strip("'")
        else:
            sheet_name = workbook.sheetnames[0]
            cell_range = item
        cell_range = cell_range.strip().strip("'").replace("$", "")
        if sheet_name not in workbook.sheetnames:
            yield sheet_name, cell_range, None
            continue
        sheet = workbook[sheet_name]
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        min_col = min_col or 1
        max_col = max_col or max(1, sheet.max_column)
        min_row = min_row or 1
        max_row = max_row or max(1, sheet.max_row)
        yield sheet_name, cell_range, (min_col, min_row, max_col, max_row)


def summarize_recalculated_targets(
    formula_workbook_path: Path,
    recalculated_workbook_path: Path,
    answer_position: str,
) -> str:
    formulas = openpyxl.load_workbook(formula_workbook_path, data_only=False, read_only=True)
    values = openpyxl.load_workbook(recalculated_workbook_path, data_only=True, read_only=True)
    try:
        lines = [
            "[ONE-SHOT LIBREOFFICE OUTPUT CHECK]",
            "LibreOffice recalculated the first genuinely modified output.xlsx.",
            "This is a target-state snapshot, not gold and not a correctness verdict.",
        ]
        for sheet_name, cell_range, bounds in _target_ranges(formulas, answer_position):
            label = f"{sheet_name}!{cell_range}"
            if bounds is None or sheet_name not in values.sheetnames:
                lines.append(f"- {label}: worksheet is missing from output.xlsx")
                continue
            min_col, min_row, max_col, max_row = bounds
            row_count = max_row - min_row + 1
            column_count = max_col - min_col + 1
            total = row_count * column_count
            indices = _sample_indices(total)
            blanks = 0
            errors = 0
            samples: list[str] = []
            formula_sheet = formulas[sheet_name]
            value_sheet = values[sheet_name]
            for index in indices:
                row = min_row + index // column_count
                column = min_col + index % column_count
                coordinate = f"{get_column_letter(column)}{row}"
                formula = formula_sheet.cell(row=row, column=column).value
                value = value_sheet.cell(row=row, column=column).value
                blanks += value in (None, "")
                errors += isinstance(value, str) and value.upper() in _FORMULA_ERRORS
                if isinstance(formula, str) and formula.startswith("="):
                    samples.append(
                        f"{coordinate}: formula={_display(formula)}, recalculated={_display(value)}"
                    )
                else:
                    samples.append(f"{coordinate}: value={_display(value)}")
            lines.append(
                f"- {label}: inspected {len(indices)}/{total}; "
                f"blank={blanks}; formula_error={errors}"
            )
            lines.extend(f"  {sample}" for sample in samples)
        lines.append(
            "Correct concrete formula errors, missing target values, or an obviously wrong target "
            "shape before completion. If the snapshot is consistent with the task, continue normally."
        )
        return "\n".join(lines)
    finally:
        formulas.close()
        values.close()


def build_libreoffice_output_feedback(
    *, output_path: Path, answer_position: str, instance_id: str
) -> str:
    from sb_adapter.evaluate import _recalculate_workbook

    with tempfile.TemporaryDirectory(prefix="v034_output_feedback_") as temporary:
        recalculated = Path(
            _recalculate_workbook(
                str(output_path),
                temporary,
                instance_id,
                timeout_seconds=180,
            )
        )
        return summarize_recalculated_targets(output_path, recalculated, answer_position)


def wrap_bash_with_output_feedback(
    bash_tool: Tool,
    *,
    working_dir: str | Path,
    answer_position: str,
    instance_id: str,
    feedback_builder: Callable[..., str] = build_libreoffice_output_feedback,
) -> Tool:
    root = Path(working_dir)
    input_path = root / "input.xlsx"
    output_path = root / "output.xlsx"
    feedback_sent = False

    def execute(command: str) -> str:
        nonlocal feedback_sent
        observation = bash_tool.execute(command=command)
        if feedback_sent or not output_path.is_file() or output_path.is_symlink():
            return observation
        if input_path.is_file() and _sha256_file(output_path) == _sha256_file(input_path):
            return observation
        feedback_sent = True
        try:
            feedback = feedback_builder(
                output_path=output_path,
                answer_position=answer_position,
                instance_id=instance_id,
            )
        except Exception as exc:
            feedback = (
                "[ONE-SHOT LIBREOFFICE OUTPUT CHECK]\n"
                f"The host recalculation check failed with {type(exc).__name__}; "
                "continue solving and validate output.xlsx with available tools."
            )
        return f"{observation}\n\n{feedback}"

    return Tool(
        name=bash_tool.name,
        description=bash_tool.description,
        func=execute,
        parameters=list(bash_tool.parameters),
    )
