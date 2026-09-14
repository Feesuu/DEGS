from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path, PurePosixPath
import re
from typing import Any, Mapping, cast
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

import openpyxl
from openpyxl.utils.cell import (
    column_index_from_string,
    get_column_letter,
    range_boundaries,
)
from openpyxl.utils.exceptions import InvalidFileException

from .core import canonical_json_bytes


TARGET_EVIDENCE_POLICY = "spreadsheetbench_target_evidence_card_v3"
WORKFLOW_CONTEXT_DOCUMENT_FORMAT = (
    "query_plus_symmetric_input_workbook_role_signature_v2"
)
STRUCTURE_ROWS = 8


class TargetContextUnavailable(ValueError):
    """The declared input cannot be observed without changing its identity."""


@dataclass(frozen=True)
class TargetEvidenceCard:
    status: str
    observations: tuple[Mapping[str, str], ...]
    input_sha256: str | None
    input_relative_path: str | None
    failure: str | None = None

    @property
    def evidence_ids(self) -> tuple[str, ...]:
        return ("Q", *(row["evidence_id"] for row in self.observations))

    def llm_payload(self) -> dict[str, Any]:
        return {
            "input_observations": [
                {
                    "evidence_id": row["evidence_id"],
                    "kind": row["kind"],
                    "value": json.loads(row["content"]),
                }
                for row in self.observations
            ]
        }

    def audit_payload(self) -> dict[str, Any]:
        payload = self.llm_payload()
        return {
            "policy": TARGET_EVIDENCE_POLICY,
            "status": self.status,
            "observations": [
                {
                    "evidence_id": row["evidence_id"],
                    "kind": row["kind"],
                    "content_sha256": hashlib.sha256(
                        row["content"].encode("utf-8")
                    ).hexdigest(),
                }
                for row in self.observations
            ],
            "input_relative_path": self.input_relative_path,
            "input_sha256": self.input_sha256,
            "card_payload_sha256": hashlib.sha256(
                canonical_json_bytes(payload)
            ).hexdigest(),
            "failure": self.failure,
        }

    def retrieval_clarification(self, evidence_ids: tuple[str, ...]) -> str:
        by_id = {row["evidence_id"]: row for row in self.observations}
        selected = []
        for evidence_id in evidence_ids:
            if evidence_id == "Q":
                continue
            try:
                selected.append(_retrieval_fact(by_id[evidence_id]))
            except KeyError as exc:
                raise ValueError("retrieval evidence ID is unavailable") from exc
        if not selected:
            raise ValueError("retrieval clarification requires input evidence")
        return "Input-only disambiguating evidence: " + " ".join(selected)

    def workflow_context_document(
        self,
        instruction: str,
        *,
        evidence_ids: tuple[str, ...] | None = None,
    ) -> str:
        """Render a symmetric source/target document for workflow context only."""
        if type(instruction) is not str or not instruction.strip():
            raise ValueError("workflow context requires a non-empty instruction")
        if self.status != "OK":
            return instruction
        selected_ids = (
            {row["evidence_id"] for row in self.observations}
            if evidence_ids is None
            else {item for item in evidence_ids if item != "Q"}
        )
        available = {row["evidence_id"] for row in self.observations}
        if not selected_ids or not selected_ids.issubset(available):
            raise ValueError("workflow context evidence differs")
        facts = [
            _workflow_role_fact(row, instruction=instruction)
            for row in self.observations
            if row["evidence_id"] in selected_ids
        ]
        return (
            f"Task instruction: {instruction}\n"
            "Input-workbook role signature:\n- "
            + "\n- ".join(facts)
        )


def _read_input_workbook(
    dataset_path: Path, spreadsheet_path: str
) -> tuple[str, bytes]:
    root = dataset_path.expanduser().resolve(strict=True).parent
    declared = PurePosixPath(spreadsheet_path)
    if (
        not spreadsheet_path
        or declared.is_absolute()
        or not declared.parts
        or any(part in {"", ".", ".."} for part in declared.parts)
    ):
        raise ValueError("spreadsheet_path must stay beneath the dataset root")
    task_dir = root.joinpath(*declared.parts)
    try:
        resolved_dir = task_dir.resolve(strict=True)
    except OSError as exc:
        raise TargetContextUnavailable("declared task directory is unavailable") from exc
    if not resolved_dir.is_relative_to(root) or not resolved_dir.is_dir():
        raise ValueError("spreadsheet_path must name a directory beneath the dataset root")
    candidates = []
    for path in sorted(resolved_dir.iterdir()):
        name = path.name.casefold()
        if not (name.endswith("_init.xlsx") or name == "initial.xlsx"):
            continue
        if any(token in name for token in ("answer", "golden", "output", "verifier")):
            continue
        candidates.append(path)
    if not candidates:
        raise TargetContextUnavailable("declared input workbook is unavailable")
    if len(candidates) != 1:
        raise ValueError("task must contain exactly one declared input workbook")
    input_path = candidates[0]
    if not input_path.is_file():
        raise ValueError("input workbook must be a regular file")
    try:
        input_bytes = input_path.read_bytes()
    except OSError as exc:
        raise TargetContextUnavailable("input workbook could not be read") from exc
    relative_path = PurePosixPath(*declared.parts, input_path.name).as_posix()
    return relative_path, input_bytes


def _cell_value(value: Any) -> Any:
    if value is None or type(value) in (str, bool, int):
        return value
    if type(value) is float:
        return value if math.isfinite(value) else str(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    value_type = type(value)
    return f"<{value_type.__module__}.{value_type.__qualname__}>"


def _observation(evidence_id: str, kind: str, value: Any) -> dict[str, str]:
    return {
        "evidence_id": evidence_id,
        "kind": kind,
        "content": json.dumps(
            value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
        ),
    }


def _format_class(value: Any) -> str:
    text = str(value).casefold()
    if any(token in text for token in ("yy", "dd", "mm/", "m/", "h:", "ss")):
        return "date_or_time"
    if "%" in text:
        return "percentage"
    if any(token in text for token in ("$", "€", "£", "¥")):
        return "currency"
    if text in {"general", "@"}:
        return "general_or_text"
    return "numeric_or_other"


def _cell_kind(value: Any) -> str:
    return {
        "f": "formula",
        "s": "text",
        "str": "text",
        "n": "numeric_or_blank",
        "d": "date_or_time",
        "b": "boolean",
        "e": "error",
    }.get(str(value), "other")


def _retrieval_fact(observation: Mapping[str, str]) -> str:
    kind = observation["kind"]
    value = json.loads(observation["content"])
    if kind == "workbook_sheet_inventory":
        return f"The input workbook contains {len(value['sheet_names'])} worksheet(s)."
    if kind == "workbook_sheet_dimensions":
        rows = [int(item["max_row"]) for item in value]
        columns = [int(item["max_column"]) for item in value]
        return (
            "Observed input-sheet extent ranges from "
            f"{min(rows)} to {max(rows)} rows and {min(columns)} to "
            f"{max(columns)} columns."
        )
    if kind == "relevant_sheet_structure":
        selected_cells = [
            cell
            for row in value["rows"]
            for cell in row["selected_cells"]
        ]
        cell_kinds = sorted({_cell_kind(cell["data_type"]) for cell in selected_cells})
        formats = sorted(
            {_format_class(cell["number_format"]) for cell in selected_cells}
        )
        return (
            "Relevant input structure has "
            f"{len(value['rows'])} observed nonempty row pattern(s); selected "
            f"cells have value kinds {cell_kinds or ['unobserved']} and format "
            f"classes {formats or ['unobserved']}."
        )
    if kind == "target_range_landmarks":
        cells = value["cells"]
        cell_kinds = sorted({_cell_kind(cell["data_type"]) for cell in cells})
        formats = sorted({_format_class(cell["number_format"]) for cell in cells})
        populated = sum(cell["value"] is not None for cell in cells)
        return (
            "The input target-neighborhood contains "
            f"{populated} populated landmark cell(s), with value kinds "
            f"{cell_kinds or ['unobserved']} and format classes "
            f"{formats or ['unobserved']}."
        )
    raise ValueError("unsupported target-evidence observation kind")


def _instruction_names_value(instruction: str, value: Any) -> bool:
    if type(value) is not str or not value.strip():
        return False
    text = value.strip().casefold()
    return re.search(
        r"(?<!\w)" + re.escape(text) + r"(?!\w)",
        instruction.casefold(),
    ) is not None


def _workflow_role_cell(
    cell: Mapping[str, Any], *, keep_text: bool, instruction: str
) -> tuple[str, str, int, int]:
    kind = _cell_kind(cell.get("data_type"))
    format_class = _format_class(cell.get("number_format"))
    value = cell.get("value")
    if (
        keep_text
        and kind == "text"
        and _instruction_names_value(instruction, value)
    ):
        role = json.dumps(str(value), ensure_ascii=False)
    elif value is None:
        role = "<blank>"
    else:
        role = f"<{kind}>"
    coordinate = str(cell["coordinate"])
    match = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", coordinate)
    if match is None:
        raise ValueError("workflow-context cell coordinate differs")
    return (
        coordinate,
        f"{role}[{format_class}]",
        int(match.group(2)),
        column_index_from_string(match.group(1)),
    )


def _workflow_role_cells(
    cells: list[Mapping[str, Any]], *, keep_text: bool, instruction: str
) -> str:
    """Run-length encode adjacent cells without dropping any role boundary."""
    encoded = [
        _workflow_role_cell(
            cell,
            keep_text=keep_text,
            instruction=instruction,
        )
        for cell in cells
    ]
    if not encoded:
        return "none"
    rendered: list[str] = []
    start_coordinate, descriptor, row, previous_column = encoded[0]
    end_coordinate = start_coordinate
    for coordinate, next_descriptor, next_row, column in encoded[1:]:
        if (
            next_descriptor == descriptor
            and next_row == row
            and column == previous_column + 1
        ):
            end_coordinate = coordinate
            previous_column = column
            continue
        coordinate_range = (
            start_coordinate
            if start_coordinate == end_coordinate
            else f"{start_coordinate}:{end_coordinate}"
        )
        rendered.append(f"{coordinate_range}={descriptor}")
        start_coordinate = end_coordinate = coordinate
        descriptor = next_descriptor
        row = next_row
        previous_column = column
    coordinate_range = (
        start_coordinate
        if start_coordinate == end_coordinate
        else f"{start_coordinate}:{end_coordinate}"
    )
    rendered.append(f"{coordinate_range}={descriptor}")
    return ",".join(rendered)


def _workflow_role_fact(observation: Mapping[str, str], *, instruction: str) -> str:
    """Keep layout/role evidence while masking non-header cell values."""
    kind = observation["kind"]
    value = json.loads(observation["content"])
    if kind == "workbook_sheet_inventory":
        return (
            f"active sheet {json.dumps(value['active_sheet'], ensure_ascii=False)}; "
            "sheet order "
            + json.dumps(value["sheet_names"], ensure_ascii=False)
        )
    if kind == "workbook_sheet_dimensions":
        return "sheet extents " + "; ".join(
            f"{json.dumps(item['sheet'], ensure_ascii=False)}="
            f"{item['used_range']}({int(item['max_row'])}x{int(item['max_column'])})"
            for item in value
        )
    if kind == "relevant_sheet_structure":
        rows = value["rows"]
        first_row = min((int(row["row"]) for row in rows), default=-1)
        rendered_rows = []
        for row in rows:
            row_number = int(row["row"])
            cells = _workflow_role_cells(
                row["selected_cells"],
                keep_text=row_number == first_row,
                instruction=instruction,
            )
            spans = ",".join(row["nonempty_column_spans"]) or "none"
            rendered_rows.append(
                f"row {row_number} spans {spans}; cells {cells}"
            )
        return (
            f"sheet {json.dumps(value['sheet'], ensure_ascii=False)} structure "
            + "; ".join(rendered_rows)
        )
    if kind == "target_range_landmarks":
        cells = _workflow_role_cells(
            value["cells"], keep_text=False, instruction=instruction
        )
        merged = ",".join(value["merged_ranges"]) or "none"
        return (
            f"target sheet {json.dumps(value['sheet'], ensure_ascii=False)} "
            f"landmarks {cells}; merged {merged}"
        )
    raise ValueError("unsupported workflow-context observation kind")


def _answer_range(
    answer_position: str, *, active_sheet: str
) -> tuple[str, tuple[int, int, int, int]] | None:
    value = answer_position.strip()
    if not value:
        return None
    if "!" in value:
        sheet, cells = value.rsplit("!", 1)
        sheet = sheet.strip()
        if len(sheet) >= 2 and sheet[0] == sheet[-1] == "'":
            sheet = sheet[1:-1].replace("''", "'")
    else:
        sheet, cells = active_sheet, value
    cells = cells.replace("$", "").strip()
    try:
        bounds = range_boundaries(cells)
    except (TypeError, ValueError):
        return None
    if any(type(bound) is not int for bound in bounds):
        return None
    return sheet, cast(tuple[int, int, int, int], bounds)


def _mentioned_sheets(instruction: str, sheet_names: list[str]) -> set[str]:
    lowered = instruction.casefold()
    return {
        name
        for name in sheet_names
        if re.search(r"(?<!\w)" + re.escape(name.casefold()) + r"(?!\w)", lowered)
    }


def _cell_payload(cell: Any) -> dict[str, Any]:
    value = "<formula>" if str(cell.data_type) == "f" else _cell_value(cell.value)
    return {
        "coordinate": cell.coordinate,
        "value": value,
        "data_type": str(cell.data_type),
        "number_format": str(cell.number_format),
    }


def _referenced_columns(
    instruction: str,
    *,
    target_bounds: tuple[int, int, int, int] | None,
    max_column: int,
) -> set[int]:
    labels = {
        match.group(1).upper()
        for match in re.finditer(
            r"(?i)(?<![A-Z0-9_])\$?([A-Z]{1,3})\$?\d+", instruction
        )
    }
    labels.update(
        match.group(1).upper()
        for match in re.finditer(
            r"(?i)\b(?:column|col\.?)\s+\$?([A-Z]{1,3})\b", instruction
        )
    )
    columns: set[int] = set()
    for label in labels:
        try:
            column = column_index_from_string(label)
        except ValueError:
            continue
        if column <= max_column:
            columns.add(column)
    if target_bounds is not None:
        columns.update(range(target_bounds[0], target_bounds[2] + 1))
    return {
        neighbor
        for column in columns
        for neighbor in (column - 1, column, column + 1)
        if 1 <= neighbor <= max_column
    }


def _column_spans(columns: list[int]) -> list[str]:
    if not columns:
        return []
    spans: list[str] = []
    start = previous = columns[0]
    for column in columns[1:]:
        if column == previous + 1:
            previous = column
            continue
        first = get_column_letter(start)
        last = get_column_letter(previous)
        spans.append(first if first == last else f"{first}:{last}")
        start = previous = column
    first = get_column_letter(start)
    last = get_column_letter(previous)
    spans.append(first if first == last else f"{first}:{last}")
    return spans


def _structure_rows(
    worksheet: Any,
    *,
    instruction: str,
    target_bounds: tuple[int, int, int, int] | None,
) -> list[dict[str, Any]]:
    selected_columns = _referenced_columns(
        instruction,
        target_bounds=target_bounds,
        max_column=worksheet.max_column,
    )
    instruction_folded = instruction.casefold()
    rows: list[dict[str, Any]] = []
    for row_index, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=min(STRUCTURE_ROWS, worksheet.max_row)),
        1,
    ):
        nonempty = [cell for cell in row if cell.value is not None]
        if not nonempty:
            continue
        cells = [
            _cell_payload(cell)
            for cell in nonempty
            if cell.column in selected_columns
            or (
                row_index == 1
                and type(cell.value) is str
                and len(cell.value.strip()) >= 2
                and cell.value.strip().casefold() in instruction_folded
            )
        ]
        rows.append(
            {
                "row": row_index,
                "nonempty_column_spans": _column_spans(
                    sorted(cell.column for cell in nonempty)
                ),
                "selected_cells": cells,
            }
        )
    return rows


def _target_landmarks(
    worksheet: Any, bounds: tuple[int, int, int, int]
) -> list[dict[str, Any]]:
    min_col, min_row, max_col, max_row = bounds
    row_indices = sorted(
        {
            row
            for row in (min_row - 1, min_row, min_row + 1, max_row - 1, max_row, max_row + 1)
            if 1 <= row <= worksheet.max_row
        }
    )
    column_indices = sorted(
        {
            column
            for column in (
                min_col - 1,
                min_col,
                min_col + 1,
                max_col - 1,
                max_col,
                max_col + 1,
            )
            if 1 <= column <= worksheet.max_column
        }
    )
    return [
        _cell_payload(worksheet.cell(row=row, column=column))
        for row in row_indices
        for column in column_indices
        if worksheet.cell(row=row, column=column).value is not None
        or min_row <= row <= max_row
        and min_col <= column <= max_col
    ]


def build_target_evidence_card(
    *,
    dataset_path: Path,
    spreadsheet_path: str,
    answer_position: str,
    instruction: str,
) -> TargetEvidenceCard:
    if any(type(value) is not str for value in (spreadsheet_path, answer_position, instruction)):
        raise ValueError("target-evidence fields must be strings")
    input_relative_path, input_bytes = _read_input_workbook(
        dataset_path, spreadsheet_path
    )
    try:
        workbook = openpyxl.load_workbook(BytesIO(input_bytes), data_only=False)
        try:
            active = workbook.active
            if active is None:
                raise TargetContextUnavailable("active worksheet is unavailable")
            observations: list[Mapping[str, str]] = []

            def add(kind: str, value: Any) -> None:
                observations.append(_observation(f"O{len(observations)}", kind, value))

            add(
                "workbook_sheet_inventory",
                {"active_sheet": active.title, "sheet_names": workbook.sheetnames},
            )
            add(
                "workbook_sheet_dimensions",
                [
                    {
                        "sheet": worksheet.title,
                        "used_range": worksheet.calculate_dimension(),
                        "max_row": worksheet.max_row,
                        "max_column": worksheet.max_column,
                    }
                    for worksheet in workbook.worksheets
                ],
            )
            parsed_target = _answer_range(answer_position, active_sheet=active.title)
            relevant = {active.title}
            relevant.update(_mentioned_sheets(instruction, workbook.sheetnames))
            if parsed_target is not None:
                relevant.add(parsed_target[0])
            for sheet_name in workbook.sheetnames:
                if sheet_name not in relevant:
                    continue
                sheet_target_bounds = (
                    parsed_target[1]
                    if parsed_target is not None and parsed_target[0] == sheet_name
                    else None
                )
                add(
                    "relevant_sheet_structure",
                    {
                        "sheet": sheet_name,
                        "rows": _structure_rows(
                            workbook[sheet_name],
                            instruction=instruction,
                            target_bounds=sheet_target_bounds,
                        ),
                    },
                )
            if parsed_target is not None and parsed_target[0] in workbook.sheetnames:
                target_sheet, bounds = parsed_target
                worksheet = workbook[target_sheet]
                add(
                    "target_range_landmarks",
                    {
                        "sheet": target_sheet,
                        "cells": _target_landmarks(worksheet, bounds),
                        "merged_ranges": [
                            str(cell_range)
                            for cell_range in worksheet.merged_cells.ranges
                            if not (
                                cell_range.max_col < bounds[0]
                                or cell_range.min_col > bounds[2]
                                or cell_range.max_row < bounds[1]
                                or cell_range.min_row > bounds[3]
                            )
                        ],
                    },
                )
        finally:
            workbook.close()
    except TargetContextUnavailable:
        raise
    except (
        BadZipFile,
        EOFError,
        InvalidFileException,
        KeyError,
        OSError,
        ParseError,
        TypeError,
        ValueError,
    ) as exc:
        raise TargetContextUnavailable(
            "verified input workbook could not be parsed"
        ) from exc
    return TargetEvidenceCard(
        status="OK",
        observations=tuple(observations),
        input_sha256=hashlib.sha256(input_bytes).hexdigest(),
        input_relative_path=input_relative_path,
    )


def unavailable_target_evidence_card(
    exc: TargetContextUnavailable,
) -> TargetEvidenceCard:
    return TargetEvidenceCard(
        status="UNAVAILABLE",
        observations=(),
        input_sha256=None,
        input_relative_path=None,
        failure=str(exc),
    )


__all__ = [
    "TARGET_EVIDENCE_POLICY",
    "WORKFLOW_CONTEXT_DOCUMENT_FORMAT",
    "TargetContextUnavailable",
    "TargetEvidenceCard",
    "build_target_evidence_card",
    "unavailable_target_evidence_card",
]
