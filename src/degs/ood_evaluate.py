"""Score DEGS table-QA OOD outputs with the official dataset implementations."""

from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import importlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from typing import Any, Mapping, Sequence

from openpyxl import load_workbook

from .core import canonical_json_bytes
from .ood_dataset import (
    ANSWER_SHEET,
    SOURCE_SPECS,
    require_clean_source_repo,
    verify_population,
)


FORMAT = "degs_tableqa_ood_official_evaluation_v1"
RUN_RESULT_FORMAT = "degs_tableqa_ood_case_result_v1"
RUN_COMPLETION_FORMAT = "degs_tableqa_ood_agent_completion_v1"


def _sha(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _read_object(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if type(value) is not dict:
        raise ValueError(f"expected JSON object: {path}")
    return value


def _write_atomic(path: Path, value: Any) -> None:
    payload = json.dumps(value, ensure_ascii=False, indent=2).encode("utf-8") + b"\n"
    staging = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    path.parent.mkdir(parents=True, exist_ok=True)
    with staging.open("xb") as handle:
        handle.write(payload)
    os.replace(staging, path)


def parse_prediction(value: Any) -> list[Any]:
    if value is None:
        return []
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return [value]
    text = str(value).strip()
    if not text:
        return []
    try:
        parsed = json.loads(text)
        return parsed if isinstance(parsed, list) else [parsed]
    except (json.JSONDecodeError, TypeError):
        pass
    if "|" in text:
        return [item.strip() for item in text.split("|")]
    if "\n" in text:
        return [item.strip() for item in text.splitlines() if item.strip()]
    return [text]


def _prediction_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    text = (
        json.dumps(value, ensure_ascii=False, separators=(",", ":"))
        if isinstance(value, (dict, list))
        else str(value)
    )
    return text.replace("\t", " ").replace("\r", " ").replace("\n", " ")


def _verify_completed_run(
    *, population: Mapping[str, Any], run_dir: Path
) -> tuple[dict[str, Any], dict[str, Any], list[dict[str, Any]]]:
    root = run_dir.expanduser().resolve()
    manifest = _read_object(root / "run_manifest.json")
    completion = _read_object(root / "results.json")
    ledger_bytes = (root / "results.jsonl").read_bytes()
    lines = ledger_bytes.splitlines()
    rows = [json.loads(line) for line in lines]
    if (
        completion.get("format") != RUN_COMPLETION_FORMAT
        or completion.get("protocol_sha256") != manifest.get("protocol_sha256")
        or completion.get("task_denominator") != population["task_count"]
        or completion.get("completed_tasks") != population["task_count"]
        or completion.get("results_jsonl_sha256") != _sha(ledger_bytes)
        or len(rows) != population["task_count"]
        or ledger_bytes != b"".join(canonical_json_bytes(row) + b"\n" for row in rows)
    ):
        raise ValueError("completed OOD run identity differs")
    for index, (task, row) in enumerate(zip(population["tasks"], rows, strict=True)):
        output_path = root / "outputs" / task["spreadsheet_path"] / task["output_file"]
        if output_path.is_symlink() or not output_path.is_file():
            actual_sha, actual_size = None, 0
        else:
            payload = output_path.read_bytes()
            actual_sha, actual_size = _sha(payload), len(payload)
        case_path = root / "case_results" / f"{_sha(str(task['task_id']).encode('utf-8'))}.json"
        case = _read_object(case_path)
        if (
            type(row) is not dict
            or row != case
            or row.get("format") != RUN_RESULT_FORMAT
            or row.get("protocol_sha256") != manifest.get("protocol_sha256")
            or row.get("dataset") != population["dataset"]
            or row.get("task_id") != task["task_id"]
            or row.get("source_id") != task["source_id"]
            or row.get("query_index") != index
            or row.get("output_path") != str(output_path)
            or row.get("output_sha256") != actual_sha
            or row.get("output_size") != actual_size
        ):
            raise ValueError(f"completed OOD task {task['task_id']} differs")
    return manifest, completion, rows


def _output_projection_sha(
    population: Mapping[str, Any], rows: Sequence[Mapping[str, Any]]
) -> str:
    return _sha(
        canonical_json_bytes(
            [
                {
                    "task_id": task["task_id"],
                    "output_sha256": row["output_sha256"],
                    "output_size": row["output_size"],
                }
                for task, row in zip(population["tasks"], rows, strict=True)
            ]
        )
    )


def _predictions(
    *, population: Mapping[str, Any], run_dir: Path
) -> tuple[dict[str, list[Any]], dict[str, str], Counter[str]]:
    predictions: dict[str, list[Any]] = {}
    errors: dict[str, str] = {}
    failures: Counter[str] = Counter()
    for task in population["tasks"]:
        task_id = str(task["task_id"])
        source_id = str(task["source_id"])
        case_path = run_dir / "case_results" / f"{_sha(task_id.encode('utf-8'))}.json"
        if not case_path.is_file():
            failures["missing_case_result"] += 1
            predictions[source_id] = []
            errors[source_id] = "case result missing"
            continue
        case = _read_object(case_path)
        failure_kind = str(case.get("failure_kind", ""))
        if failure_kind:
            failures[failure_kind] += 1
        output_path = run_dir / "outputs" / task["spreadsheet_path"] / task["output_file"]
        if output_path.is_symlink() or not output_path.is_file():
            failures["missing_output"] += 1
            predictions[source_id] = []
            errors[source_id] = "output workbook missing"
            continue
        try:
            workbook = load_workbook(output_path, data_only=False, read_only=True)
            try:
                predictions[source_id] = parse_prediction(
                    workbook[ANSWER_SHEET].cell(row=1, column=2).value
                )
            finally:
                workbook.close()
            errors[source_id] = ""
        except Exception as exc:
            failures["unreadable_output"] += 1
            predictions[source_id] = []
            errors[source_id] = f"{type(exc).__name__}: {exc}"
    if len(predictions) != population["task_count"]:
        raise ValueError("OOD prediction source-ID coverage differs")
    return predictions, errors, failures


def _score_wikitq(
    *,
    population: Mapping[str, Any],
    predictions: Mapping[str, list[Any]],
    errors: Mapping[str, str],
    source_repo: Path,
    output_dir: Path,
    python2: str,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    lines = []
    for task in population["tasks"]:
        source_id = str(task["source_id"])
        lines.append(
            "\t".join(
                [source_id, *(_prediction_text(item) for item in predictions[source_id])]
            )
        )
    prediction_path = output_dir / "wikitq_predictions.tsv"
    prediction_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    command = [
        python2,
        str(source_repo / "evaluator.py"),
        "-t",
        str(source_repo / "tagged/data"),
        str(prediction_path),
    ]
    completed = subprocess.run(command, check=True, capture_output=True)
    stdout = completed.stdout.decode("utf-8", errors="replace")
    stderr = completed.stderr.decode("utf-8", errors="replace")
    (output_dir / "wikitq_official_stdout.tsv").write_text(stdout, encoding="utf-8")
    (output_dir / "wikitq_official_stderr.log").write_text(stderr, encoding="utf-8")
    scores: dict[str, bool] = {}
    for line in stdout.splitlines():
        fields = line.split("\t", 2)
        if len(fields) >= 2 and fields[1] in {"True", "False"}:
            scores[fields[0]] = fields[1] == "True"
    if len(scores) != population["task_count"]:
        raise RuntimeError(f"WikiTQ official evaluator coverage differs: {len(scores)}")
    rows = [
        {
            "task_id": task["task_id"],
            "source_id": task["source_id"],
            "prediction": predictions[str(task["source_id"])],
            "official_passed": scores[str(task["source_id"])],
            "output_error": errors[str(task["source_id"])],
        }
        for task in population["tasks"]
    ]
    passed = sum(row["official_passed"] for row in rows)
    return rows, {
        "dataset": "wikitq",
        "evaluator": "official evaluator.py with tagged/data",
        "denominator": len(rows),
        "passed": passed,
        "accuracy": passed / len(rows),
        "official_stderr_tail": stderr.strip().splitlines()[-5:],
    }


def _score_hitab(
    *,
    population: Mapping[str, Any],
    predictions: Mapping[str, list[Any]],
    errors: Mapping[str, str],
    source_repo: Path,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    samples = [
        json.loads(line)
        for line in (source_repo / "data/test_samples.jsonl").read_text(encoding="utf-8").splitlines()
        if line
    ]
    gold = {str(row["id"]): row["answer"] for row in samples}
    if len(gold) != population["task_count"]:
        raise ValueError("HiTab official gold population differs")
    sys.path.insert(0, str(source_repo))
    try:
        hmt_score = importlib.import_module("qa.table.utils").hmt_score
    finally:
        sys.path.pop(0)
    rows = []
    score_errors = 0
    for task in population["tasks"]:
        source_id = str(task["source_id"])
        prediction = predictions[source_id]
        score_error = ""
        if not prediction:
            passed = False
        else:
            try:
                passed = bool(hmt_score(prediction, gold[source_id]))
            except Exception as exc:
                passed = False
                score_error = f"{type(exc).__name__}: {exc}"
                score_errors += 1
        rows.append(
            {
                "task_id": task["task_id"],
                "source_id": source_id,
                "prediction": prediction,
                "official_passed": passed,
                "output_error": errors[source_id],
                "official_score_error": score_error,
            }
        )
    passed = sum(row["official_passed"] for row in rows)
    return rows, {
        "dataset": "hitab",
        "evaluator": "official qa.table.utils.hmt_score",
        "denominator": len(rows),
        "passed": passed,
        "accuracy": passed / len(rows),
        "official_score_errors": score_errors,
    }


def verify_evaluation(
    *,
    prepared_data_path: Path,
    run_dir: Path,
    source_repo: Path,
    output_dir: Path,
) -> dict[str, Any]:
    population = verify_population(prepared_data_path)
    resolved_run_dir = run_dir.expanduser().resolve()
    run_manifest, completion, run_rows = _verify_completed_run(
        population=population,
        run_dir=resolved_run_dir,
    )
    source_repo = source_repo.expanduser().resolve()
    spec = SOURCE_SPECS[population["dataset"]]
    require_clean_source_repo(source_repo, expected_commit=str(spec["commit"]))
    predictions, errors, failures = _predictions(
        population=population,
        run_dir=resolved_run_dir,
    )
    output_dir = output_dir.expanduser().resolve()
    summary_path = output_dir / "eval_summary.json"
    scores_path = output_dir / "official_scores.jsonl"
    if (
        output_dir.is_symlink()
        or not output_dir.is_dir()
        or summary_path.is_symlink()
        or scores_path.is_symlink()
        or not summary_path.is_file()
        or not scores_path.is_file()
    ):
        raise FileNotFoundError("OOD evaluation artifact differs")
    summary = _read_object(summary_path)
    rows_bytes = scores_path.read_bytes()
    lines = rows_bytes.splitlines()
    rows = [json.loads(line) for line in lines]
    unsigned = {key: value for key, value in summary.items() if key != "self_sha256"}
    passed = sum(type(row) is dict and row.get("official_passed") is True for row in rows)
    row_fields = (
        {"task_id", "source_id", "prediction", "official_passed", "output_error"}
        if population["dataset"] == "wikitq"
        else {
            "task_id",
            "source_id",
            "prediction",
            "official_passed",
            "output_error",
            "official_score_error",
        }
    )
    ordered_rows = all(
        type(row) is dict
        and set(row) == row_fields
        and row.get("task_id") == task["task_id"]
        and row.get("source_id") == task["source_id"]
        and row.get("prediction") == predictions[str(task["source_id"])]
        and row.get("output_error") == errors[str(task["source_id"])]
        and type(row.get("official_passed")) is bool
        for row, task in zip(rows, population["tasks"], strict=True)
    ) if len(rows) == population["task_count"] else False
    if (
        summary.get("format") != FORMAT
        or summary.get("self_sha256") != _sha(canonical_json_bytes(unsigned))
        or summary.get("dataset") != population["dataset"]
        or summary.get("source_commit") != spec["commit"]
        or summary.get("population_manifest_sha256") != population["self_sha256"]
        or summary.get("run_protocol_sha256") != run_manifest.get("protocol_sha256")
        or summary.get("run_results_jsonl_sha256")
        != completion.get("results_jsonl_sha256")
        or summary.get("run_output_projection_sha256")
        != _output_projection_sha(population, run_rows)
        or summary.get("failure_counts") != dict(sorted(failures.items()))
        or summary.get("denominator") != population["task_count"]
        or summary.get("passed") != passed
        or summary.get("accuracy") != passed / population["task_count"]
        or summary.get("official_scores_sha256") != _sha(rows_bytes)
        or len(lines) != population["task_count"]
        or not ordered_rows
        or any(canonical_json_bytes(row) != line for row, line in zip(rows, lines, strict=True))
    ):
        raise ValueError("OOD evaluation identity differs")
    return summary


def evaluate(
    *,
    prepared_data_path: Path,
    run_dir: Path,
    source_repo: Path,
    output_dir: Path,
    python2: str = "/usr/bin/python2",
) -> dict[str, Any]:
    population = verify_population(prepared_data_path)
    run_dir = run_dir.expanduser().resolve()
    run_manifest, completion, run_rows = _verify_completed_run(
        population=population,
        run_dir=run_dir,
    )
    source_repo = source_repo.expanduser().resolve()
    output_dir = output_dir.expanduser().absolute()
    spec = SOURCE_SPECS[population["dataset"]]
    require_clean_source_repo(source_repo, expected_commit=str(spec["commit"]))
    if output_dir.exists() or output_dir.is_symlink():
        return verify_evaluation(
            prepared_data_path=prepared_data_path,
            run_dir=run_dir,
            source_repo=source_repo,
            output_dir=output_dir,
        )
    output_dir.parent.mkdir(parents=True, mode=0o700, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=f".{output_dir.name}.", dir=output_dir.parent))
    published = False
    try:
        predictions, errors, failures = _predictions(population=population, run_dir=run_dir)
        if population["dataset"] == "wikitq":
            rows, summary = _score_wikitq(
                population=population,
                predictions=predictions,
                errors=errors,
                source_repo=source_repo,
                output_dir=staging,
                python2=python2,
            )
        else:
            rows, summary = _score_hitab(
                population=population,
                predictions=predictions,
                errors=errors,
                source_repo=source_repo,
            )
        rows_bytes = b"".join(canonical_json_bytes(row) + b"\n" for row in rows)
        (staging / "official_scores.jsonl").write_bytes(rows_bytes)
        body = {
            "format": FORMAT,
            **summary,
            "source_commit": spec["commit"],
            "population_manifest_sha256": population["self_sha256"],
            "run_protocol_sha256": run_manifest["protocol_sha256"],
            "run_results_jsonl_sha256": completion["results_jsonl_sha256"],
            "run_output_projection_sha256": _output_projection_sha(
                population, run_rows
            ),
            "official_scores_sha256": _sha(rows_bytes),
            "failure_counts": dict(sorted(failures.items())),
        }
        result = {**body, "self_sha256": _sha(canonical_json_bytes(body))}
        _write_atomic(staging / "eval_summary.json", result)
        os.replace(staging, output_dir)
        published = True
    finally:
        if not published:
            shutil.rmtree(staging, ignore_errors=True)
    verify_evaluation(
        prepared_data_path=prepared_data_path,
        run_dir=run_dir,
        source_repo=source_repo,
        output_dir=output_dir,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return result


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--prepared-data-path", type=Path, required=True)
    parser.add_argument("--run-dir", type=Path, required=True)
    parser.add_argument("--source-repo", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--python2", default="/usr/bin/python2")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    evaluate(
        prepared_data_path=args.prepared_data_path,
        run_dir=args.run_dir,
        source_repo=args.source_repo,
        output_dir=args.output_dir,
        python2=args.python2,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
