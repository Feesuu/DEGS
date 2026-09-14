"""Prepare and verify input-only WikiTQ/HiTab workbook populations."""

from __future__ import annotations

import argparse
from io import BytesIO
import csv
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from typing import Any, Mapping, Sequence
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .core import canonical_json_bytes


FORMAT = "degs_tableqa_ood_input_population_v1"
ANSWER_SHEET = "Answer"
ANSWER_CELL = "B1"
ANSWER_POSITION = f"{ANSWER_SHEET}!{ANSWER_CELL}"
SOURCE_SPECS = {
    "wikitq": {
        "commit": "7d455a5a707b96341ef72aff9428749d443d8aa9",
        "count": 2810,
        "url": "https://github.com/ppasupat/WikiTableQuestions.git",
    },
    "hitab": {
        "commit": "d179602662b490249baf068a76fbe4137029126e",
        "count": 1584,
        "url": "https://github.com/microsoft/HiTab.git",
    },
}
_DATASET_FIELDS = {
    "id",
    "dataset",
    "source_id",
    "instruction",
    "question",
    "instruction_type",
    "answer_position",
    "spreadsheet_path",
}
_TASK_FIELDS = {
    "query_index",
    "dataset_index",
    "task_id",
    "dataset",
    "source_id",
    "semantic_query",
    "instruction",
    "instruction_type",
    "answer_position",
    "spreadsheet_path",
    "input_file",
    "output_file",
    "input_sha256",
    "input_size",
}


def _sha(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _safe_relative(value: Any, *, label: str) -> Path:
    if type(value) is not str or not value:
        raise ValueError(f"{label} must be a non-empty relative path")
    path = Path(value)
    if path.is_absolute() or any(part in {"", ".", ".."} for part in path.parts):
        raise ValueError(f"{label} must be a safe relative path")
    return path


def _write_exclusive(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, mode=0o700, exist_ok=True)
    descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(descriptor, "wb") as handle:
        handle.write(payload)


def _repo_head(repo: Path) -> str:
    completed = subprocess.run(
        ["git", "-C", str(repo), "rev-parse", "HEAD"],
        check=True,
        capture_output=True,
        text=True,
    )
    return completed.stdout.strip()


def require_clean_source_repo(repo: Path, *, expected_commit: str) -> None:
    """Bind preparation/scoring to the checked-in official dataset state."""

    if _repo_head(repo) != expected_commit:
        raise ValueError("OOD source repository commit differs")
    status = subprocess.run(
        ["git", "-C", str(repo), "status", "--porcelain", "--untracked-files=all"],
        check=True,
        capture_output=True,
        text=True,
    ).stdout
    if status.strip():
        raise ValueError("OOD source repository has local changes")


def _safe_id(prefix: str, raw_id: str) -> str:
    return f"{prefix}_{re.sub(r'[^A-Za-z0-9_.-]+', '_', str(raw_id))}"


def _write_literal(cell: Any, value: Any) -> None:
    if value is None:
        cell.value = ""
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        cell.value = value
    else:
        cell.value = str(value)
        if str(value).startswith("="):
            cell.data_type = "s"


def _style_workbook(workbook: Workbook, *, hierarchical: bool = False) -> None:
    sheet = workbook["Table"]
    if not hierarchical:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for column in range(1, min(sheet.max_column, 40) + 1):
        width = 12
        for row in range(1, min(sheet.max_row, 100) + 1):
            width = max(width, min(len(str(sheet.cell(row, column).value or "")) + 2, 40))
        sheet.column_dimensions[get_column_letter(column)].width = width
    answer = workbook.create_sheet(ANSWER_SHEET)
    answer["A1"] = "Final answer"
    answer["A1"].font = Font(bold=True)
    answer[ANSWER_CELL] = ""
    answer.column_dimensions["A"].width = 18
    answer.column_dimensions["B"].width = 60


def _workbook_bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _qa_instruction(question: str) -> str:
    return (
        "Answer the question using the provided spreadsheet. "
        f"Write only the final answer to {ANSWER_POSITION}.\n\nQuestion: {question}"
    )


def _tsv_unescape(value: str) -> str:
    return value.replace(r"\n", "\n").replace(r"\p", "|").replace("\\\\", "\\")


def _leaf_paths(root: Mapping[str, Any]) -> list[list[str]]:
    indexed: dict[int, list[str]] = {}

    def visit(node: Mapping[str, Any], path: list[str]) -> None:
        name = str(node.get("name", node.get("value", "")))
        current = path if name.startswith("<") else [*path, name]
        if node.get("line_idx") is not None:
            indexed[int(node["line_idx"])] = current or [""]
        for child in node.get("children_dict") or []:
            visit(child, current)

    visit(root, [])
    return [path for _, path in sorted(indexed.items())]


def _build_hitab_workbook(table: Mapping[str, Any]) -> bytes:
    top_paths = _leaf_paths(table["top_root"])
    left_paths = _leaf_paths(table["left_root"])
    data = table["data"]
    if len(data) != len(left_paths):
        raise ValueError("HiTab row hierarchy differs from its data matrix")
    if data and len(data[0]) != len(top_paths):
        raise ValueError("HiTab column hierarchy differs from its data matrix")
    top_depth = max((len(path) for path in top_paths), default=1)
    left_depth = max((len(path) for path in left_paths), default=1)
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("new workbook has no active worksheet")
    sheet.title = "Table"
    _write_literal(sheet.cell(1, 1), table.get("title", ""))
    for level in range(top_depth):
        for column, path in enumerate(top_paths, start=left_depth + 1):
            padded = [""] * (top_depth - len(path)) + path
            _write_literal(sheet.cell(2 + level, column), padded[level])
    for row_index, path in enumerate(left_paths, start=2 + top_depth):
        padded = [""] * (left_depth - len(path)) + path
        for column, value in enumerate(padded, start=1):
            _write_literal(sheet.cell(row_index, column), value)
        for column, item in enumerate(data[row_index - (2 + top_depth)], start=left_depth + 1):
            _write_literal(sheet.cell(row_index, column), item.get("value"))
    _style_workbook(workbook, hierarchical=True)
    sheet.freeze_panes = sheet.cell(2 + top_depth, left_depth + 1).coordinate
    return _workbook_bytes(workbook)


def _wikitq_rows(repo: Path, *, expected_count: int) -> list[tuple[str, str, str, bytes]]:
    split = repo / "data/random-split-1-dev.tsv"
    with split.open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle, delimiter="\t"))
    if len(rows) != expected_count:
        raise ValueError(f"WikiTQ population differs: {len(rows)} != {expected_count}")
    table_cache: dict[str, bytes] = {}
    result = []
    for row in rows:
        context = row["context"]
        if context not in table_cache:
            with (repo / context).open(encoding="utf-8", newline="") as handle:
                table_rows = list(csv.reader(handle))
            workbook = Workbook()
            sheet = workbook.active
            if sheet is None:
                raise RuntimeError("new workbook has no active worksheet")
            sheet.title = "Table"
            for row_index, table_row in enumerate(table_rows, start=1):
                for column_index, value in enumerate(table_row, start=1):
                    _write_literal(sheet.cell(row_index, column_index), value)
            _style_workbook(workbook)
            table_cache[context] = _workbook_bytes(workbook)
        question = _tsv_unescape(row["utterance"])
        result.append((str(row["id"]), _safe_id("wikitq", row["id"]), question, table_cache[context]))
    return result


def _hitab_rows(repo: Path, *, expected_count: int) -> list[tuple[str, str, str, bytes]]:
    samples = [
        json.loads(line)
        for line in (repo / "data/test_samples.jsonl").read_text(encoding="utf-8").splitlines()
        if line
    ]
    if len(samples) != expected_count:
        raise ValueError(f"HiTab population differs: {len(samples)} != {expected_count}")
    table_cache: dict[str, bytes] = {}
    result = []
    with zipfile.ZipFile(repo / "data/tables.zip") as archive:
        for sample in samples:
            table_id = str(sample["table_id"])
            if table_id not in table_cache:
                table = json.loads(archive.read(f"tables/hmt/{table_id}.json"))
                table_cache[table_id] = _build_hitab_workbook(table)
            source_id = str(sample["id"])
            result.append(
                (
                    source_id,
                    _safe_id("hitab", source_id),
                    str(sample["question"]),
                    table_cache[table_id],
                )
            )
    return result


def _publish_population(
    *,
    dataset: str,
    rows: Sequence[tuple[str, str, str, bytes]],
    output: Path,
) -> dict[str, Any]:
    tasks: list[dict[str, Any]] = []
    dataset_rows: list[dict[str, Any]] = []
    for index, (source_id, task_id, question, workbook_bytes) in enumerate(rows):
        spreadsheet_path = f"spreadsheet/{task_id}"
        input_file = f"1_{task_id}_init.xlsx"
        output_file = f"1_{task_id}_output.xlsx"
        _write_exclusive(output / spreadsheet_path / input_file, workbook_bytes)
        instruction = _qa_instruction(question)
        dataset_rows.append(
            {
                "id": task_id,
                "dataset": dataset,
                "source_id": source_id,
                "instruction": instruction,
                "question": question,
                "instruction_type": (
                    "hierarchical_table_question_answering"
                    if dataset == "hitab"
                    else "table_question_answering"
                ),
                "answer_position": ANSWER_POSITION,
                "spreadsheet_path": spreadsheet_path,
            }
        )
        tasks.append(
            {
                "query_index": index,
                "dataset_index": index,
                "task_id": task_id,
                "dataset": dataset,
                "source_id": source_id,
                "semantic_query": question,
                "instruction": instruction,
                "instruction_type": dataset_rows[-1]["instruction_type"],
                "answer_position": ANSWER_POSITION,
                "spreadsheet_path": spreadsheet_path,
                "input_file": input_file,
                "output_file": output_file,
                "input_sha256": _sha(workbook_bytes),
                "input_size": len(workbook_bytes),
            }
        )
    dataset_bytes = json.dumps(dataset_rows, ensure_ascii=False, indent=2).encode("utf-8") + b"\n"
    _write_exclusive(output / "dataset.json", dataset_bytes)
    input_projection = [
        {
            "task_id": row["task_id"],
            "input_sha256": row["input_sha256"],
            "input_size": row["input_size"],
        }
        for row in tasks
    ]
    body = {
        "format": FORMAT,
        "dataset": dataset,
        "source_url": SOURCE_SPECS[dataset]["url"],
        "source_commit": SOURCE_SPECS[dataset]["commit"],
        "task_count": len(tasks),
        "dataset_json_sha256": _sha(dataset_bytes),
        "query_projection_sha256": _sha(
            canonical_json_bytes(
                [
                    {
                        "query_index": row["query_index"],
                        "task_id": row["task_id"],
                        "semantic_query": row["semantic_query"],
                    }
                    for row in tasks
                ]
            )
        ),
        "input_tree_sha256": _sha(canonical_json_bytes(input_projection)),
        "tasks": tasks,
        "gold_available": False,
    }
    manifest = {**body, "self_sha256": _sha(canonical_json_bytes(body))}
    _write_exclusive(output / "retrieval_manifest.json", canonical_json_bytes(manifest))
    return manifest


def prepare_population(*, dataset: str, source_repo: Path, output_dir: Path) -> dict[str, Any]:
    if dataset not in SOURCE_SPECS:
        raise ValueError(f"unsupported OOD dataset: {dataset}")
    repo = source_repo.expanduser().resolve()
    expected = SOURCE_SPECS[dataset]
    require_clean_source_repo(repo, expected_commit=str(expected["commit"]))
    output = output_dir.expanduser().absolute()
    if output.exists() or output.is_symlink():
        raise FileExistsError("OOD prepared population directory must be fresh")
    output.parent.mkdir(parents=True, mode=0o700, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=f".{output.name}.", dir=output.parent))
    published = False
    try:
        rows = (
            _wikitq_rows(repo, expected_count=int(expected["count"]))
            if dataset == "wikitq"
            else _hitab_rows(repo, expected_count=int(expected["count"]))
        )
        manifest = _publish_population(
            dataset=dataset,
            rows=rows,
            output=staging,
        )
        os.replace(staging, output)
        published = True
    finally:
        if not published:
            shutil.rmtree(staging, ignore_errors=True)
    return verify_population(output)


def verify_population(root: Path) -> dict[str, Any]:
    base = root.expanduser().absolute()
    if base.is_symlink() or not base.is_dir():
        raise FileNotFoundError("OOD prepared population differs")
    base = base.resolve()
    manifest_path = base / "retrieval_manifest.json"
    dataset_path = base / "dataset.json"
    if any(path.is_symlink() or not path.is_file() for path in (manifest_path, dataset_path)):
        raise FileNotFoundError("OOD prepared population manifest differs")
    manifest_bytes = manifest_path.read_bytes()
    manifest = json.loads(manifest_bytes)
    if type(manifest) is not dict:
        raise ValueError("OOD retrieval manifest differs")
    unsigned = {key: value for key, value in manifest.items() if key != "self_sha256"}
    dataset = manifest.get("dataset")
    tasks = manifest.get("tasks")
    expected = SOURCE_SPECS.get(dataset) if type(dataset) is str else None
    dataset_bytes = dataset_path.read_bytes()
    dataset_rows = json.loads(dataset_bytes)
    if (
        manifest_bytes != canonical_json_bytes(manifest)
        or manifest.get("format") != FORMAT
        or manifest.get("self_sha256") != _sha(canonical_json_bytes(unsigned))
        or expected is None
        or manifest.get("source_commit") != expected["commit"]
        or manifest.get("source_url") != expected["url"]
        or manifest.get("task_count") != expected["count"]
        or manifest.get("gold_available") is not False
        or manifest.get("dataset_json_sha256") != _sha(dataset_bytes)
        or type(tasks) is not list
        or len(tasks) != expected["count"]
        or type(dataset_rows) is not list
        or len(dataset_rows) != expected["count"]
    ):
        raise ValueError("OOD prepared population identity differs")
    task_ids: set[str] = set()
    source_ids: set[str] = set()
    expected_files = {"dataset.json", "retrieval_manifest.json"}
    input_projection = []
    query_projection = []
    for index, (task, dataset_row) in enumerate(zip(tasks, dataset_rows, strict=True)):
        if (
            type(task) is not dict
            or set(task) != _TASK_FIELDS
            or type(dataset_row) is not dict
            or set(dataset_row) != _DATASET_FIELDS
            or task.get("query_index") != index
            or task.get("dataset_index") != index
            or task.get("dataset") != dataset
            or dataset_row.get("dataset") != dataset
            or task.get("task_id") != dataset_row.get("id")
            or task.get("source_id") != dataset_row.get("source_id")
            or task.get("semantic_query") != dataset_row.get("question")
            or task.get("instruction") != dataset_row.get("instruction")
            or task.get("instruction_type") != dataset_row.get("instruction_type")
            or task.get("answer_position") != ANSWER_POSITION
            or dataset_row.get("answer_position") != ANSWER_POSITION
            or task.get("spreadsheet_path") != dataset_row.get("spreadsheet_path")
            or task["task_id"] in task_ids
            or task["source_id"] in source_ids
        ):
            raise ValueError(f"OOD prepared task {index} differs")
        if any("gold" in str(key).casefold() for key in (*task.keys(), *dataset_row.keys())):
            raise ValueError("OOD input-only population contains a gold field")
        spreadsheet_path = _safe_relative(task["spreadsheet_path"], label="spreadsheet path")
        input_file = _safe_relative(task["input_file"], label="input file")
        output_file = _safe_relative(task["output_file"], label="output file")
        expected_spreadsheet_path = Path("spreadsheet") / str(task["task_id"])
        if (
            spreadsheet_path != expected_spreadsheet_path
            or len(input_file.parts) != 1
            or len(output_file.parts) != 1
            or input_file.name != f"1_{task['task_id']}_init.xlsx"
            or output_file.name != f"1_{task['task_id']}_output.xlsx"
        ):
            raise ValueError(f"OOD prepared workbook routing {index} differs")
        input_path = base / spreadsheet_path / input_file
        expected_files.add(input_path.relative_to(base).as_posix())
        payload = input_path.read_bytes()
        if (
            input_path.is_symlink()
            or not input_path.is_file()
            or task.get("input_sha256") != _sha(payload)
            or task.get("input_size") != len(payload)
        ):
            raise ValueError(f"OOD prepared workbook {index} differs")
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False)
        try:
            if "Table" not in workbook.sheetnames or ANSWER_SHEET not in workbook.sheetnames:
                raise ValueError(f"OOD prepared workbook {index} has wrong sheets")
            if workbook[ANSWER_SHEET][ANSWER_CELL].value not in {None, ""}:
                raise ValueError("OOD input workbook contains an answer")
        finally:
            workbook.close()
        task_ids.add(task["task_id"])
        source_ids.add(task["source_id"])
        input_projection.append(
            {
                "task_id": task["task_id"],
                "input_sha256": task["input_sha256"],
                "input_size": task["input_size"],
            }
        )
        query_projection.append(
            {
                "query_index": index,
                "task_id": task["task_id"],
                "semantic_query": task["semantic_query"],
            }
        )
    actual_files: set[str] = set()
    for path in base.rglob("*"):
        if path.is_symlink():
            raise ValueError("OOD prepared population contains a symlink")
        if path.is_file():
            actual_files.add(path.relative_to(base).as_posix())
        elif not path.is_dir():
            raise ValueError("OOD prepared population contains a special file")
    if (
        actual_files != expected_files
        or manifest.get("input_tree_sha256") != _sha(canonical_json_bytes(input_projection))
        or manifest.get("query_projection_sha256") != _sha(canonical_json_bytes(query_projection))
    ):
        raise ValueError("OOD prepared population projection differs")
    return manifest


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    prepare = sub.add_parser("prepare")
    prepare.add_argument("--dataset", choices=sorted(SOURCE_SPECS), required=True)
    prepare.add_argument("--source-repo", type=Path, required=True)
    prepare.add_argument("--output-dir", type=Path, required=True)
    verify = sub.add_parser("verify")
    verify.add_argument("--output-dir", type=Path, required=True)
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    manifest = (
        prepare_population(
            dataset=args.dataset,
            source_repo=args.source_repo,
            output_dir=args.output_dir,
        )
        if args.command == "prepare"
        else verify_population(args.output_dir)
    )
    print(
        json.dumps(
            {
                "dataset": manifest["dataset"],
                "task_count": manifest["task_count"],
                "self_sha256": manifest["self_sha256"],
            },
            sort_keys=True,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
